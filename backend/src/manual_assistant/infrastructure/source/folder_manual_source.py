"""Leitura da pasta de manuais da Engenharia (no servidor, uma pasta de rede montada).

SOMENTE LEITURA: este módulo só lista, consulta e lê arquivos. Não existe nele nenhuma
operação de escrita, e assim deve continuar: a pasta pertence à Engenharia.
"""

import asyncio
import io
import logging
from collections.abc import Callable, Mapping, Sequence
from pathlib import Path, PurePosixPath
from typing import TypeVar

from openpyxl import load_workbook

from manual_assistant.application.errors import ExternalServiceError
from manual_assistant.application.ports.manual_source import SourceDocument

logger = logging.getLogger(__name__)

# Unidade de rede costuma falhar na primeira leitura depois de parada (reconexão):
# tenta de novo antes de desistir. Erros "definitivos" não são repetidos.
ATTEMPTS = 3
_DEFINITIVE_ERRORS = (FileNotFoundError, NotADirectoryError, PermissionError)

T = TypeVar("T")


class FolderManualSource:
    """Estrutura esperada (a da pasta da Engenharia):

        <raiz>/
          MANUAIS PRODUTO.xlsm      planilha com a descrição de cada código (opcional)
          ARMAZENAGEM/*.pdf         manuais vigentes
          PROTEÍNA/*.pdf
          PROTEÍNA/Versões Anteriores/...   ignoradas: só o 1º nível de cada subpasta conta

    PDFs soltos na raiz também são ignorados.
    """

    def __init__(
        self, root: Path, *, catalog_file: str | None = None, retry_delay_seconds: float = 2
    ) -> None:
        self._root = root
        self._catalog_file = catalog_file
        self._retry_delay_seconds = retry_delay_seconds

    async def list_documents(self) -> Sequence[SourceDocument]:
        try:
            files = await self._with_retries(self._list_pdfs)
        except OSError as error:
            raise ExternalServiceError(
                f"Não foi possível ler a pasta de manuais ({self._root}): {error}"
            ) from error
        titles = await asyncio.to_thread(self._read_catalog)
        return [
            SourceDocument(
                path=relative,
                file_name=name,
                fingerprint=fingerprint,
                title=titles.get(PurePosixPath(name).stem.upper()),
            )
            for relative, name, fingerprint in files
        ]

    async def read(self, path: str) -> bytes:
        file = self._resolve(path)
        try:
            return await self._with_retries(file.read_bytes)
        except OSError as error:
            raise ExternalServiceError(f"Não foi possível ler {path}: {error}") from error

    async def _with_retries(self, operation: Callable[[], T]) -> T:
        for attempt in range(1, ATTEMPTS):
            try:
                return await asyncio.to_thread(operation)
            except _DEFINITIVE_ERRORS:
                raise
            except OSError as error:
                logger.warning("Falha ao ler %s (%s); tentando de novo", self._root, error)
                await asyncio.sleep(self._retry_delay_seconds * attempt)
        return await asyncio.to_thread(operation)  # última tentativa: o erro sobe

    def _list_pdfs(self) -> list[tuple[str, str, str]]:
        found = []
        for folder in sorted(p for p in self._root.iterdir() if p.is_dir()):
            for file in sorted(folder.iterdir()):
                # "~$..." são arquivos temporários do Office; não são manuais.
                if file.suffix.lower() != ".pdf" or file.name.startswith("~$"):
                    continue
                if not file.is_file():
                    continue
                stat = file.stat()
                relative = file.relative_to(self._root).as_posix()
                found.append((relative, file.name, f"{stat.st_size}-{stat.st_mtime_ns}"))
        return found

    def _read_catalog(self) -> Mapping[str, str]:
        """Código (ex.: "95007003-01P") -> descrição. Sem planilha, os títulos ficam
        com o nome do arquivo: a sincronização não depende dela."""
        if not self._catalog_file:
            return {}
        path = self._root / self._catalog_file
        try:
            # Lê os bytes e abre a cópia em memória: o arquivo da rede não fica aberto.
            content = path.read_bytes()
            return _parse_catalog(content)
        except FileNotFoundError:
            logger.warning("Planilha de manuais não encontrada: %s", path)
        except Exception:
            # Planilha com formato inesperado não pode impedir a sincronização.
            logger.warning("Não foi possível ler a planilha de manuais %s", path, exc_info=True)
        return {}

    def _resolve(self, path: str) -> Path:
        root = self._root.resolve()
        file = (root / path).resolve()
        # Só arquivos de dentro da pasta (nada de "../") e só PDFs.
        if not file.is_relative_to(root) or file.suffix.lower() != ".pdf":
            raise ValueError(f"Caminho fora da pasta de manuais: {path!r}")
        return file


def _parse_catalog(content: bytes) -> dict[str, str]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    titles: dict[str, str] = {}
    try:
        for sheet in workbook.worksheets:
            columns: tuple[int, int] | None = None
            for row in sheet.iter_rows(values_only=True):
                cells = [str(value).strip() if value is not None else "" for value in row]
                if columns is None:
                    columns = _header_columns(cells)
                    continue
                code_column, description_column = columns
                if max(columns) >= len(cells):
                    continue
                code, description = cells[code_column].upper(), cells[description_column]
                if code and description:
                    titles[code] = description
    finally:
        workbook.close()
    return titles


def _header_columns(cells: list[str]) -> tuple[int, int] | None:
    """Posição das colunas "Código" e "Descrição" na linha de cabeçalho, se for ela."""
    normalized = [cell.casefold() for cell in cells]
    if "código" in normalized and "descrição" in normalized:
        return normalized.index("código"), normalized.index("descrição")
    return None
